"""Shared per-object inputs and conservative import of approved category tables.

The legacy storage key barrier_assessment is retained to preserve existing
projects and attachment references. It does not limit inputs to one work mode.
"""
from __future__ import annotations

import copy
import hashlib
import math
import re
import shutil
from pathlib import Path

from .barriers import CATEGORIES, prepare_rooms
from .io_utils import utc_now, write_json


FIRE_PUMP_SYSTEMS = (("aup", "АУП"), ("vpv", "ВПВ"))


def fire_pump_defaults():
    return {key: {"present": None, "locations": "", "source": ""}
            for key, _label in FIRE_PUMP_SYSTEMS}


def normalize_fire_pump_rooms(value):
    """Keep unanswered questions distinct from a confirmed absence."""
    if not isinstance(value, dict):
        raise ValueError("Насосные АУП/ВПВ: требуется ответ по каждой системе")
    result = fire_pump_defaults()
    for key, label in FIRE_PUMP_SYSTEMS:
        answer = value.get(key, {})
        if not isinstance(answer, dict):
            raise ValueError(f"Насосные {label}: некорректный ответ")
        result[key].update(copy.deepcopy(answer))
        present = result[key]["present"]
        if present is not None and type(present) is not bool:
            raise ValueError(f"Насосные {label}: выберите Есть, Нет или Не установлено")
        for field in ("locations", "source"):
            if not isinstance(result[key][field], str):
                raise ValueError(f"Насосные {label}: помещения и источник должны быть текстом")
            result[key][field] = result[key][field].strip()
        if present is False and result[key]["locations"]:
            raise ValueError(f"Насосные {label}: выбран ответ «Нет», но указаны помещения. Уточните ответ")
    return result


def fire_pump_rooms_markdown(project):
    data = inputs_for(project)
    free_text = " ".join(data.get("fire_pump_rooms_text", "").split())
    lines = ["## Насосные АУП и ВПВ — ответы опросника", ""]
    if free_text:
        lines.extend([
            f"- Сведения пользователя: {free_text}",
            "",
            "Сопоставь текст с планами и паспортом. Установи, какие помещения являются насосными АУП, ВПВ "
            "или общими, их расположение и источник сведений. Если текст не даёт однозначного ответа, запроси уточнение; "
            "не считай обычную хозяйственную насосную противопожарной только по названию. Проверь самостоятельные "
            "требования к ограждениям, дверям и перекрытиям по применимым нормам в рамках выбранного вида проверки."
        ])
        return "\n".join(lines)
    answers = data["fire_pump_rooms"]
    for key, label in FIRE_PUMP_SYSTEMS:
        answer = answers[key]
        status = {True: "Есть", False: "Нет", None: "Не установлено"}[answer["present"]]
        # A location may have multiple lines in an imported/manual answer.
        location = " ".join(answer["locations"].split()) or "не указаны"
        source = " ".join(answer["source"].split()) or "не указан"
        lines.append(f"- {label}: {status}. Помещения / этажи / оси: {location}. Источник: {source}.")
        if answer["present"] is None:
            lines.append(f"  Требуется уточнить, какие насосные обслуживают {label}; отсутствие ответа не означает отсутствие насосной.")
        elif answer["present"] is True and not answer["locations"]:
            lines.append(f"  Требуется адресно уточнить расположение насосных {label}.")
    lines.extend(["", "Сопоставь ответы с планами и паспортом; при противоречии запроси уточнение. "
                  "Одно помещение может обслуживать обе системы. Подтверждение относится только к указанным помещениям. "
                  "Не считай любую насосную противопожарной только по её названию. "
                  "Проверь самостоятельные требования к её ограждениям, дверям и перекрытиям по применимым нормам в рамках выбранного вида проверки."])
    return "\n".join(lines)


def defaults():
    return {"categories_approved": False, "category_approval_source": "", "category_source": {},
            "rooms": [], "stu_status": "unknown", "stu_documents": [],
            "refrigeration_present": None, "refrigeration_equipment_rooms_confirmed": None,
            "refrigeration_equipment_rooms": [], "additional_documents_reviewed": False,
            "preparation_completed": False, "source_explanations_reviewed": False,
            "room_match_completed": False, "attachments": [], "special_conditions": "",
            "lifts_and_hoists": "", "corridor_division": "", "slabs_and_openings": "",
            "storage_and_equipment": "", "fire_pump_rooms_text": "",
            "fire_pump_rooms": fire_pump_defaults()}


def inputs_for(project):
    result = defaults()
    result.update(copy.deepcopy(project.get("barrier_assessment", {})))
    result["fire_pump_rooms"] = normalize_fire_pump_rooms(result["fire_pump_rooms"])
    return result


def additional_inputs_markdown(project):
    data = inputs_for(project)
    return (
        "## Дополнительные данные объекта — для всех видов проверок\n\n"
        "Прочитай данные вкладки «Дополнительные данные» вместе с паспортом перед работой в любом режиме. "
        "Они относятся только к выбранному объекту, а не только к оценке преград.\n\n"
        f"- Помещений в таблице: {len(data.get('rooms', []))}.\n"
        f"- Дополнительных документов: {len(data.get('attachments', []))}.\n"
        "- Проверь таблицу помещений, категорий и источник её утверждения; статус и документы СТУ; "
        "другие вложения и ответы заказчика; холодильное оборудование и насосные; лифты, подъёмники, "
        "коридоры, перекрытия, заполнения проёмов, хранение и остальные уточнения.\n\n"
        "Полные значения, пути и контрольные суммы вложений приведены ниже в полном снимке объекта. "
        "Раздел barrier_assessment имеет историческое служебное имя: его содержимое является "
        "общими дополнительными данными объекта для всех видов проверок. Не игнорируй его из-за имени. "
        "Открой относящиеся к задаче файлы, а не ограничивайся именами вложений. При недоступности "
        "материала зафиксируй недостаток данных; наличие пути не означает, что документ прочитан. "
        "Применимость каждого сведения устанавливай в рамках выбранной проверки. Отметки готовности "
        "не заменяют нормативный анализ; отсутствие данных не означает отсутствие оборудования или требований. "
        "Доступность этих сведений не разрешает применять алгоритм ПП в других режимах или начинать "
        "вид проверки, алгоритм которого ещё не утверждён.\n"
    )


def fire_pump_text_for_ui(data):
    """Return the free answer, or preserve meaningful answers from the former structured form."""
    text = data.get("fire_pump_rooms_text", "")
    if isinstance(text, str) and text.strip():
        return text.strip()
    answers = normalize_fire_pump_rooms(data.get("fire_pump_rooms", {}))
    parts = []
    for key, label in FIRE_PUMP_SYSTEMS:
        answer = answers[key]
        if answer["present"] is None and not answer["locations"] and not answer["source"]:
            continue
        status = {True: "есть", False: "нет", None: "наличие не установлено"}[answer["present"]]
        details = [f"Насосная {label}: {status}"]
        if answer["locations"]:
            details.append(f"помещения / этажи / оси: {answer['locations']}")
        if answer["source"]:
            details.append(f"источник: {answer['source']}")
        parts.append(", ".join(details))
    return "; ".join(parts)


def finite_number(value, label, *, optional=True, positive=True):
    if value is None or value == "":
        if optional:
            return None
        raise ValueError(f"Не заполнено: {label}")
    if isinstance(value, bool):
        raise ValueError(f"{label}: требуется число")
    value = float(str(value).replace(",", ".").replace(" ", ""))
    if not math.isfinite(value) or (positive and value <= 0):
        raise ValueError(f"{label}: требуется положительное конечное число")
    return value


def validate_inputs(data):
    if data.get("stu_status") not in {"unknown", "absent", "provided"}:
        raise ValueError("СТУ: выберите неизвестно, отсутствуют или представлены")
    rooms = prepare_rooms(data.get("rooms", []))
    for room in rooms:
        room["area_m2"] = finite_number(room.get("area_m2"), "Площадь помещения")
        room["storage_height_m"] = finite_number(room.get("storage_height_m"), "Высота хранения")
        if room.get("host_part_functional_class", "") not in {"", "Ф4.3", "Ф5.1", "Ф5.2", "другое"}:
            raise ValueError("Неизвестный функциональный класс части здания")
    result = copy.deepcopy(data)
    result["rooms"] = rooms
    pump_text = result.get("fire_pump_rooms_text", "")
    if not isinstance(pump_text, str):
        raise ValueError("Насосные АУП/ВПВ: сведения должны быть текстом")
    result["fire_pump_rooms_text"] = pump_text.strip()
    result["fire_pump_rooms"] = normalize_fire_pump_rooms(data.get("fire_pump_rooms", {}))
    if result.get("categories_approved") and not str(result.get("category_approval_source", "")).strip():
        raise ValueError("Укажите, кем/каким сообщением утверждены категории")
    if result.get("stu_status") == "absent" and result.get("stu_documents"):
        raise ValueError("Указаны документы СТУ, но выбран статус «Отсутствуют»")
    return result


def room_key(row):
    return str(row.get("floor", "")).strip() + ":" + str(row.get("number", "")).strip()


def category_values(rows):
    return [{k:r.get(k) for k in ("floor", "number", "name", "category", "area_m2")} for r in rows]


def snapshot_manual_categories(paths, slug, data):
    """A manual table has its own immutable version, retaining its Excel origin."""
    import json
    values = category_values(data["rooms"])
    payload = {"rooms": values, "origin": data.get("category_source", {})}
    content = json.dumps(payload, ensure_ascii=False, sort_keys=True, indent=2).encode("utf-8")
    sha = hashlib.sha256(content).hexdigest()
    folder = paths.projects / slug / "sources" / "attachments"
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / (sha[:12] + "_категории_ручной_ввод.json")
    if not target.exists():
        target.write_bytes(content)
    return {"path":str(target.resolve()), "sha256":sha, "role":"categories", "manual":True,
            "recorded_at":utc_now()}


def snapshot_attachment(paths, slug, source, role):
    path = Path(source).resolve()
    if not path.is_file():
        raise ValueError("Исходный файл не найден")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    folder = paths.projects / slug / "sources" / "attachments"
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / (digest[:12] + "_" + path.name)
    if not target.exists():
        shutil.copy2(path, target)
    return {"path": str(target.resolve()), "original_path": str(path), "sha256": digest,
            "role": role, "recorded_at": utc_now()}


def read_category_workbook(path):
    """Import values, never calculate categories or infer their approval."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    path = Path(path).resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Таблицу категорий передайте в формате .xlsx")
    wb = load_workbook(path, read_only=True, data_only=True)
    rows, skipped = [], []
    normalize = lambda x: re.sub(r"[^а-яa-z0-9]", "", str(x or "").casefold().replace("ё", "е"))
    try:
        for sheet in wb:
            mapping = None
            for line, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                if mapping is None:
                    if line > 30:
                        break
                    headers = [normalize(v) for v in cells]
                    def index(test):
                        found = [i for i, h in enumerate(headers) if test(h)]
                        return found[0] if len(found) == 1 else None
                    candidate = {
                        "number": index(lambda h: h in {"номер", "номерпомещения", "номерпомещ", "помещ", "помещения"}),
                        "name": index(lambda h: "наименование" in h or h in {"название", "помещение"}),
                        "category": index(lambda h: h.startswith("категория")),
                        "area_m2": index(lambda h: h.startswith("площадьпомещ") or h in {"площадь", "площадьм2"}),
                        "floor": index(lambda h: h == "этаж"),
                    }
                    if all(candidate[k] is not None for k in ("number", "name", "category")):
                        mapping = candidate
                    continue
                def cell(key):
                    i = mapping[key]
                    return cells[i] if i is not None and i < len(cells) else None
                number, name, category = cell("number"), cell("name"), cell("category")
                if all(v is None or str(v).strip() == "" for v in (number, name, category)):
                    continue
                if number is None and normalize(name).startswith("итог"):
                    continue
                if isinstance(number, (int, float)) and float(number).is_integer():
                    number = str(int(number))
                category = str(category or "").strip().upper()
                if category not in CATEGORIES or not str(number or "").strip() or not str(name or "").strip():
                    raise ValueError(f"{sheet.title}, строка {line}: не заполнены номер, название или утверждённая категория")
                rows.append({"floor": str(cell("floor") or sheet.title).strip(), "number": str(number).strip(),
                    "name": str(name).strip(), "category": category,
                    "area_m2": finite_number(cell("area_m2"), f"{sheet.title}:{line}, площадь"),
                    "source_sheet": sheet.title,
                    "source_range": f"{get_column_letter(min(i for i in mapping.values() if i is not None)+1)}{line}:"
                                    f"{get_column_letter(max(i for i in mapping.values() if i is not None)+1)}{line}",
                    "source_file": path.name})
            if mapping is None:
                skipped.append(sheet.title)
    finally:
        wb.close()
    if not rows:
        raise ValueError("Не найдены однозначные столбцы номера, наименования и категории. Исправьте заголовки таблицы или введите помещения вручную.")
    return {"rooms": prepare_rooms(rows), "skipped_sheets": skipped,
            "source": {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}}


def merge_category_rows(previous, imported):
    """Keep equipment/part data; changed categories invalidate drawing matches."""
    old = {room_key(r): r for r in previous}
    merged = []
    for row in imported:
        before = old.get(room_key(row), {})
        result = {**copy.deepcopy(before), **row}
        if any(before.get(k) != row.get(k) for k in ("name", "category", "area_m2")):
            for key in ("matched_cad_handles", "cad_match_status", "category_entity_handle"):
                result.pop(key, None)
        merged.append(result)
    return merged
